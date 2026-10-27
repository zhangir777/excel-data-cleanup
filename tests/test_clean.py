import zipfile
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from cleanup.clean import (
    clean_email,
    clean_name,
    clean_rows,
    map_header,
    normalize_phone,
    parse_amount,
    parse_date,
)
from cleanup.cli import main, read_table
from cleanup.report import write_report


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("1 200,50 ₸", 1200.5),
        ("$1,200.50", 1200.5),
        ("1.200,50", 1200.5),
        ("1,200", 1200.0),
        ("15500,5", 15500.5),
        ("1.000.000", 1000000.0),
        ("8000", 8000.0),
        (7000, 7000.0),
        ("free", None),
        ("", None),
    ],
)
def test_parse_amount(raw, expected):
    assert parse_amount(raw) == expected


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("03.02.2026", date(2026, 2, 3)),
        ("2026-02-03", date(2026, 2, 3)),
        ("2026-02-03 14:20", date(2026, 2, 3)),
        ("03/02/2026", date(2026, 2, 3)),
        ("03.02.26", date(2026, 2, 3)),
        ("03 Feb 2026", date(2026, 2, 3)),
        ("soon", None),
        ("", None),
    ],
)
def test_parse_date_day_first(raw, expected):
    assert parse_date(raw) == expected


def test_parse_date_month_first():
    assert parse_date("02/03/2026", month_first=True) == date(2026, 2, 3)


def test_small_helpers():
    assert normalize_phone("8 777 123 45 67") == "+77771234567"
    assert normalize_phone("123") is None
    assert clean_name("  aLIYA   nurlanovna-ova ") == "Aliya Nurlanovna-Ova"
    assert clean_email(" Aliya@Example.COM ") == "aliya@example.com"
    assert clean_email("aliya at example.com") is None
    assert map_header("Amount, KZT") == "amount" and map_header("Дата") == "date"
    assert map_header("Unrelated") is None


def row(**kw):
    base = {"Date ": "01.08.2026", "Client name": "aliya", "Phone": "87771234567",
            "E-mail": "a@example.com", "Manager": "asel", "Product": "Haircut",
            "Amount, KZT": "8 000 ₸"}  # fmt: skip
    return base | kw


def test_clean_rows_counts_and_reasons():
    rows = [
        row(),
        row(),  # duplicate
        row(**{"Phone": "+7 (777) 123-45-67"}),  # same phone in another format: duplicate too
        row(**{"Date ": "soon"}),
        row(**{"Amount, KZT": "free"}),
        {k: "" for k in row()},
        row(**{"Date ": "2026-07-15", "Phone": "123", "E-mail": "bad"}),
    ]
    result = clean_rows(rows)
    s = result.stats
    assert (s["rows_in"], s["empty"], s["invalid"], s["duplicates"], s["clean"]) == (7, 1, 2, 2, 2)
    assert s["bad_phones"] == 1 and s["bad_emails"] == 1
    assert [r["date"] for r in result.clean] == [date(2026, 7, 15), date(2026, 8, 1)]  # sorted
    first = result.clean[1]
    assert (first["client"], first["manager"], first["amount"]) == ("Aliya", "Asel", 8000.0)
    reasons = sorted(r["reason"] for r in result.removed)
    assert reasons == ["duplicate", "duplicate", "no valid amount", "no valid date"]


def test_report_and_cli(tmp_path):
    src = tmp_path / "in.csv"
    src.write_text(
        "Date ;Client name;Phone;Amount, KZT;Manager\n"
        "01.08.2026;aliya;87771234567;8 000 ₸;asel\n"
        "15.09.2026;marat;8 701 555 66 77;20000;kamila\n"
        ";;;;\n",
        encoding="utf-8-sig",
    )
    assert len(read_table(src)) == 3
    out = tmp_path / "out.xlsx"
    main([str(src), "--out", str(out)])
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Clean data", "Removed rows"]
    data = wb["Clean data"]
    assert data.max_row == 3 and data["C2"].value == "+77771234567"
    assert data.freeze_panes == "A2" and data["A2"].number_format == "DD.MM.YYYY"
    assert wb["Summary"]._charts
    with zipfile.ZipFile(out) as z:  # chart must point at the real sheet name
        assert "'Summary'!" in z.read("xl/charts/chart1.xml").decode()


def test_read_xlsx(tmp_path):
    wb = Workbook()
    wb.active.append(["Дата", "Клиент", "Сумма"])
    wb.active.append([datetime(2026, 8, 1, 10, 30), "aliya", 8000])
    wb.save(tmp_path / "in.xlsx")
    rows = read_table(tmp_path / "in.xlsx")
    assert rows == [{"Дата": datetime(2026, 8, 1, 10, 30), "Клиент": "aliya", "Сумма": 8000}]
    assert clean_rows(rows).clean[0]["date"] == date(2026, 8, 1)


def test_write_report_empty(tmp_path):
    path = write_report(clean_rows([]), tmp_path / "empty.xlsx")
    assert load_workbook(path)["Clean data"].max_row == 1
