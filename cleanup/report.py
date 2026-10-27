"""Write the cleaning result to a client-ready Excel file."""

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from cleanup.clean import COLUMNS, Result

HEADER_FILL = PatternFill("solid", fgColor="1A3A5C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLES = ["Date", "Client", "Phone", "Email", "Manager", "Product", "Amount"]
WIDTHS = [12, 26, 16, 28, 18, 22, 14]
STAT_LABELS = {
    "rows_in": "Rows in the source file",
    "empty": "Empty rows removed",
    "invalid": "Rows without date / amount / client",
    "duplicates": "Duplicates removed",
    "clean": "Clean rows",
    "phones_fixed": "Phones fixed to +7 format",
    "bad_phones": "Phones that are not valid numbers",
    "bad_emails": "Invalid emails cleared",
}


def _header(ws: Worksheet, titles: list[str], widths: list[int]) -> None:
    ws.append(titles)
    for i, cell in enumerate(ws[1]):
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = widths[i]
    ws.freeze_panes = "A2"


def _data_sheet(ws: Worksheet, rows: list[dict]) -> None:
    _header(ws, TITLES, WIDTHS)
    for r in rows:
        ws.append([r[c] or ("" if c != "amount" else 0) for c in COLUMNS])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
        row[6].number_format = "#,##0.00"
    ws.auto_filter.ref = ws.dimensions


def _removed_sheet(ws: Worksheet, rows: list[dict]) -> None:
    headers = list(dict.fromkeys(k for r in rows for k in r if k != "reason"))
    _header(ws, ["Reason", *headers], [26] + [18] * len(headers))
    for r in rows:
        ws.append([r["reason"], *(str(r.get(h) or "") for h in headers)])


def _table(ws: Worksheet, top: int, title: str, rows: list[tuple]) -> int:
    ws.cell(top, 1, title).font = Font(bold=True, size=12)
    for i, (label, value) in enumerate(rows, start=top + 1):
        ws.cell(i, 1, label)
        ws.cell(i, 2, value).number_format = "#,##0.00" if isinstance(value, float) else "General"
    return top + len(rows) + 2


def _summary_sheet(ws: Worksheet, result: Result) -> None:
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 36, 16
    ws.page_setup.orientation = "landscape"  # prints on one page with the chart
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    by_month, by_manager = defaultdict(float), defaultdict(float)
    for r in result.clean:
        by_month[r["date"].strftime("%Y-%m")] += r["amount"]
        by_manager[r["manager"] or "(no manager)"] += r["amount"]
    total = sum(r["amount"] for r in result.clean)
    row = _table(ws, 1, "What was cleaned", [(STAT_LABELS[k], v) for k, v in result.stats.items()])
    row = _table(ws, row, "Totals", [("Total amount", round(total, 2)),
                                     ("Orders", len(result.clean))])  # fmt: skip
    month_top = row
    row = _table(ws, row, "Amount by month", sorted(by_month.items()))
    _table(ws, row, "Amount by manager", sorted(by_manager.items(), key=lambda kv: -kv[1]))

    if by_month:
        chart = BarChart()
        chart.title, chart.legend = "Amount by month", None
        chart.height, chart.width = 7, 14
        chart.y_axis.scaling.min = 0  # bars start at zero, so months compare honestly
        data = Reference(ws, min_col=2, min_row=month_top + 1, max_row=month_top + len(by_month))
        cats = Reference(ws, min_col=1, min_row=month_top + 1, max_row=month_top + len(by_month))
        chart.add_data(data)
        chart.set_categories(cats)
        ws.add_chart(chart, "D2")


def write_report(result: Result, path: Path) -> Path:
    wb = Workbook()
    wb.active.title = "Summary"  # before the chart: chart references use the sheet name
    _summary_sheet(wb.active, result)
    _data_sheet(wb.create_sheet("Clean data"), result.clean)
    _removed_sheet(wb.create_sheet("Removed rows"), result.removed)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
